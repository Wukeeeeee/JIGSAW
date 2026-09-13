"""文件 → 文本 统一提取层
========================
readfile 工具和知识库检索（knowledge_search）共用这一层：
不管什么扩展名，都先提成纯文本再做后续处理。
支持：md/txt/json/csv/py 等纯文本、docx、pdf、xlsx、pptx。
提不出来的（图片等）返回 None，调用方决定跳过。
"""
import os


def extract_text(path: str) -> str | None:
    """把任意常见文件提成纯文本。失败/不支持返回 None，成功返回字符串（可能为空）。"""
    if not path or not os.path.isfile(path):
        return None
    ext = os.path.splitext(path)[1].lower()

    # ---- 纯文本类：直接按 utf-8 读 ----
    if ext in (".md", ".txt", ".markdown", ".json", ".csv", ".py", ".js", ".ts",
               ".html", ".htm", ".css", ".log", ".ini", ".yaml", ".yml", ".xml",
               ".toml", ".sql", ".sh", ".bat", ".cfg", ".conf"):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            return None

    # ---- Word ----
    if ext == ".docx":
        try:
            import docx2txt
            return docx2txt.process(path) or ""
        except Exception:
            return None

    # ---- PDF ----
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            r = PdfReader(path)
            return "\n".join(p.extract_text() or "" for p in r.pages)
        except Exception:
            return None

    # ---- Excel ----
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(rows)
        except Exception:
            return None

    # ---- PPT ----
    if ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
            return "\n".join(texts)
        except Exception:
            return None

    return None
